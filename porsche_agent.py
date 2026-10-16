#!/usr/bin/env python3
"""
Agente de sanitização de vendas Porsche
=======================================

Lê a planilha bruta (xlsx), aplica as regras do Schema_Porsche.md e grava uma
nova planilha com as colunas sanitizadas logo após cada coluna de origem.

Uso:
    python porsche_agent.py [entrada.xlsx] [schema.md] [saida.xlsx]

Padrões:
    entrada = /mnt/user-data/uploads/Tabela_Porche.xlsx
    schema  = /mnt/user-data/uploads/Schema_Porsche.md
    saida   = /mnt/user-data/outputs/Tabela_Porsche_Sanitizada.xlsx

O agente é guiado pelo schema: colunas obrigatórias, modelos canônicos,
formas de pagamento e status de entrega são lidos do arquivo .md.
"""
import re
import sys
import difflib
import datetime as dt
from copy import copy
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment

INVALID = "INVALID"
KM_TO_MI = Decimal("0.621371")

# ----------------------------------------------------------------------------
# 1. Leitura do schema
# ----------------------------------------------------------------------------

def _section(md: str, start: str, end: str | None) -> str:
    i = md.index(start)
    j = md.index(end, i) if end else len(md)
    return md[i:j]


def _bullets(block: str) -> list[str]:
    return re.findall(r"^- `([^`]+)`", block, flags=re.M)


def load_schema(path: str) -> dict:
    md = open(path, encoding="utf-8").read()
    cols_block = _section(md, "## Canonical input columns", "## Sanitized output columns")
    required = [
        m.group(1)
        for m in re.finditer(r"^\| `(\w+)` \| yes", cols_block, flags=re.M)
    ]
    out_block = _section(md, "## Sanitized output columns", "## Sanitization rules")
    outputs = re.findall(r"^\| `(\w+)` \| `(\w+)` \|", out_block, flags=re.M)  # (src, out)

    schema = {
        "required": required,
        "outputs": outputs,
        "models": _bullets(_section(md, "### Porsche models", "### Model year")),
        "payments": _bullets(_section(md, "### Payment method", "### City")),
        "deliveries": _bullets(_section(md, "### Delivery status", "## Quality checks")),
    }
    for k in ("required", "outputs", "models", "payments", "deliveries"):
        if not schema[k]:
            raise SystemExit(f"Schema inválido: seção '{k}' não encontrada.")
    return schema


# ----------------------------------------------------------------------------
# 2. Utilitários numéricos
# ----------------------------------------------------------------------------

_UNITS = {
    "zero": 0, "one": 1, "two": 2, "three": 3, "four": 4, "five": 5, "six": 6,
    "seven": 7, "eight": 8, "nine": 9, "ten": 10, "eleven": 11, "twelve": 12,
    "thirteen": 13, "fourteen": 14, "fifteen": 15, "sixteen": 16,
    "seventeen": 17, "eighteen": 18, "nineteen": 19,
}
_TENS = {
    "twenty": 20, "thirty": 30, "forty": 40, "fifty": 50, "sixty": 60,
    "seventy": 70, "eighty": 80, "ninety": 90,
}


def words_to_int(tokens: list[str]) -> int | None:
    """'eighty two thousand' -> 82000. Retorna None se não for número por extenso."""
    if not tokens:
        return None
    total, current = 0, 0
    for t in tokens:
        if t == "and":
            continue
        if t in _UNITS:
            current += _UNITS[t]
        elif t in _TENS:
            current += _TENS[t]
        elif t == "hundred":
            current = max(current, 1) * 100
        elif t == "thousand":
            total += max(current, 1) * 1000
            current = 0
        elif t == "million":
            total += max(current, 1) * 1_000_000
            current = 0
        else:
            return None
    return total + current


def _tokens(s: str) -> list[str]:
    return [t for t in re.split(r"[\s\-]+", s.lower().strip()) if t]


def parse_localized_number(s: str) -> Decimal | None:
    """
    Interpreta números com separadores ambíguos:
      '153,200.50' -> 153200.50     '89.750,00' -> 89750.00
      '112.750'    -> 112750        '1,200'     -> 1200
      '12.5'       -> 12.5
    Regra: com '.' e ',' juntos, o último é o decimal. Com um só tipo,
    separador repetido ou seguido de exatamente 3 dígitos = milhar.
    """
    s = s.strip().replace(" ", "")
    if not re.fullmatch(r"\d[\d.,]*", s):
        return None
    if "." in s and "," in s:
        dec = "." if s.rfind(".") > s.rfind(",") else ","
        thou = "," if dec == "." else "."
        s = s.replace(thou, "").replace(dec, ".")
    elif "." in s or "," in s:
        sep = "." if "." in s else ","
        parts = s.split(sep)
        if len(parts) > 2 or len(parts[-1]) == 3:
            s = "".join(parts)
        else:
            s = ".".join(parts)
    try:
        return Decimal(s)
    except Exception:
        return None


def _is_blank(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


# ----------------------------------------------------------------------------
# 3. Sanitizadores por coluna
# ----------------------------------------------------------------------------

_MONTHS = {
    m: i + 1 for i, m in enumerate(
        ["january", "february", "march", "april", "may", "june", "july",
         "august", "september", "october", "november", "december"])
}


def _month_from_name(name: str) -> int | None:
    n = name.lower().rstrip(".")
    if len(n) < 3:
        return None
    for full, num in _MONTHS.items():
        if full == n or (len(n) <= len(full) and full.startswith(n)):
            return num
    return None


def _mk_date(y: int, m: int, d: int) -> str:
    try:
        return dt.date(y, m, d).isoformat()
    except ValueError:
        return INVALID


def sanitize_date(v) -> str:
    if _is_blank(v):
        return INVALID
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()

    # YYYY-MM-DD, YYYY/MM/DD, YYYY.MM.DD
    m = re.fullmatch(r"(\d{4})([-/.])(\d{1,2})\2(\d{1,2})", s)
    if m:
        return _mk_date(int(m[1]), int(m[3]), int(m[4]))

    # MM/DD/YYYY, MM/DD/YY, MM-DD-YY, MM.DD.YY
    m = re.fullmatch(r"(\d{1,2})([-/.])(\d{1,2})\2(\d{2}|\d{4})", s)
    if m:
        y = int(m[4])
        if len(m[4]) == 2:
            y += 2000
        return _mk_date(y, int(m[1]), int(m[3]))

    # Month DDth, YYYY  |  Mon DDth YYYY
    m = re.fullmatch(r"([A-Za-z]+)\.?\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})", s, flags=re.I)
    if m:
        mon = _month_from_name(m[1])
        if mon:
            return _mk_date(int(m[3]), mon, int(m[2]))
    return INVALID


def sanitize_model_year(v) -> int | str:
    if _is_blank(v):
        return INVALID
    year = None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        year = int(v) if float(v).is_integer() else None
    else:
        s = str(v).strip()
        if re.fullmatch(r"\d{4}", s):
            year = int(s)
        elif re.fullmatch(r"(\d{2})[\s\-](\d{2})", s):          # 20-24 / 20 24
            year = int(s[:2]) * 100 + int(s[-2:])
        else:
            toks = _tokens(s)
            if toks[:2] == ["two", "thousand"]:                   # two thousand twenty one
                year = words_to_int(toks)
            else:                                                 # twenty twenty four
                for i in range(1, len(toks)):
                    a, b = words_to_int(toks[:i]), words_to_int(toks[i:])
                    if a is not None and b is not None and 10 <= a <= 99 and 0 <= b <= 99:
                        year = a * 100 + b
                        break
    if year is None or not (1990 <= year <= 2035):
        return INVALID
    return year


def sanitize_price(v) -> float | str:
    if _is_blank(v):
        return INVALID
    amount = None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        amount = Decimal(str(v))
    else:
        s = str(v).strip().lower()
        s = re.sub(r"\b(usd|us\$|dollars?)\b", " ", s).replace("$", " ").strip()
        words = words_to_int(_tokens(s)) if re.search(r"[a-z]", s) else None
        if words is not None:
            amount = Decimal(words)
        else:
            mult = Decimal(1)
            if re.search(r"\dk$", s.replace(" ", "")):
                mult = Decimal(1000)
                s = re.sub(r"k$", "", s.replace(" ", ""))
            num = parse_localized_number(s)
            if num is not None:
                amount = num * mult
    if amount is None or amount <= 0:
        return INVALID
    return float(amount.quantize(Decimal("0.01"), rounding=ROUND_HALF_UP))


def _round_int(d: Decimal) -> int:
    return int(d.quantize(Decimal(1), rounding=ROUND_HALF_UP))


def sanitize_mileage(v) -> int | str:
    if _is_blank(v):
        return INVALID
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return _round_int(Decimal(str(v))) if v >= 0 else INVALID

    s = str(v).strip().lower()
    is_km = bool(re.search(r"\bkm\b|\bkms\b|kilomet", s))
    s = re.sub(r"\b(miles?|mi|km|kms|kilometers?|kilometres?)\b\.?:?", " ", s)
    s = re.sub(r"(?<=\d)(miles?|mi|km)\b\.?", " ", s).replace(":", " ").strip()

    if s in ("new", "new car", "zero", "zero miles", "") and not re.search(r"\d", s):
        if s in ("new", "new car", "zero", "zero miles"):
            return 0
        # texto vazio depois de remover unidades, ex.: "miles" -> não seguro
        return INVALID

    toks = _tokens(s)
    if toks == ["new"] or toks == ["new", "car"]:
        return 0
    n = words_to_int(toks)
    num = Decimal(n) if n is not None else parse_localized_number(s)
    if num is None:
        return INVALID
    if is_km:
        num = num * KM_TO_MI
    return _round_int(num)


# --- Estados -----------------------------------------------------------------

US_STATES = {
    "alabama": "AL", "alaska": "AK", "arizona": "AZ", "arkansas": "AR",
    "california": "CA", "colorado": "CO", "connecticut": "CT", "delaware": "DE",
    "district of columbia": "DC", "florida": "FL", "georgia": "GA", "hawaii": "HI",
    "idaho": "ID", "illinois": "IL", "indiana": "IN", "iowa": "IA", "kansas": "KS",
    "kentucky": "KY", "louisiana": "LA", "maine": "ME", "maryland": "MD",
    "massachusetts": "MA", "michigan": "MI", "minnesota": "MN", "mississippi": "MS",
    "missouri": "MO", "montana": "MT", "nebraska": "NE", "nevada": "NV",
    "new hampshire": "NH", "new jersey": "NJ", "new mexico": "NM", "new york": "NY",
    "north carolina": "NC", "north dakota": "ND", "ohio": "OH", "oklahoma": "OK",
    "oregon": "OR", "pennsylvania": "PA", "rhode island": "RI",
    "south carolina": "SC", "south dakota": "SD", "tennessee": "TN", "texas": "TX",
    "utah": "UT", "vermont": "VT", "virginia": "VA", "washington": "WA",
    "west virginia": "WV", "wisconsin": "WI", "wyoming": "WY",
}
_ABBR = set(US_STATES.values())


def sanitize_state(v) -> str:
    if _is_blank(v):
        return INVALID
    s = re.sub(r"\s+", " ", str(v).strip().lower().replace(".", ""))
    if s.upper() in _ABBR:
        return s.upper()
    return US_STATES.get(s, INVALID)


# --- Cidade ------------------------------------------------------------------

def sanitize_city(v) -> str:
    if _is_blank(v):
        return INVALID
    s = re.sub(r"\s+", " ", str(v).strip())
    return " ".join(w[:1].upper() + w[1:].lower() for w in s.split(" "))


# --- Modelo ------------------------------------------------------------------

def _collapse(s: str) -> str:
    return re.sub(r"\s+", " ", s.strip().lower())


def _smart_title(s: str) -> str:
    out = []
    for t in re.sub(r"\s+", " ", s.strip()).split(" "):
        tl = t.lower()
        if re.fullmatch(r"\d+[a-z]*", tl) or re.fullmatch(r"gt\d*s?|rs", tl):
            out.append(t.upper())
        elif tl == "e-hybrid":
            out.append("E-Hybrid")
        else:
            out.append(t.capitalize())
    return " ".join(out)


def make_model_sanitizer(canonical: list[str]):
    lookup = {_collapse(m): m for m in canonical}

    def sanitize_model(v) -> str:
        if _is_blank(v):
            return INVALID
        key = _collapse(str(v))
        return lookup.get(key) or _smart_title(str(v))

    return sanitize_model


# --- Pagamento ---------------------------------------------------------------

PAYMENT_ALIASES = {
    "credit card": "Credit Card", "credit": "Credit Card",
    "debit card": "Debit Card", "debit": "Debit Card",
    "bank transfer": "Bank Transfer",
    "wire transfer": "Wire Transfer", "wire": "Wire Transfer", "bank wire": "Wire Transfer",
    "financing": "Financing", "finance": "Financing",
    "lease": "Lease", "leasing": "Lease",
    "cash": "Cash",
    "ach payment": "ACH Payment", "ach": "ACH Payment",
    "crypto payment": "Crypto Payment", "crypto": "Crypto Payment",
    "cryptocurrency": "Crypto Payment",
}


def _text_key(v) -> str:
    s = re.sub(r"(?<=[a-z])(?=[A-Z])", " ", str(v).strip())      # CreditCard -> Credit Card
    s = re.sub(r"[-_/]+", " ", s)
    s = re.sub(r"[^A-Za-z0-9 ]", "", s)
    return re.sub(r"\s+", " ", s).strip().lower()


def make_payment_sanitizer(canonical: list[str]):
    canon_by_key = {_text_key(c): c for c in canonical}
    for alias, target in PAYMENT_ALIASES.items():
        assert target in canonical, f"Alias '{alias}' aponta para rótulo fora do schema: {target}"
    table = {**PAYMENT_ALIASES, **canon_by_key}

    def sanitize_payment(v) -> str:
        if _is_blank(v):
            return INVALID
        key = _text_key(v)
        if key in table:
            return table[key]
        stripped = re.sub(r"\s+(payment|plan)$", "", key)
        if stripped in table:
            return table[stripped]
        return " ".join(w.capitalize() for w in key.split()) or INVALID

    return sanitize_payment


# --- Status de entrega -------------------------------------------------------

def make_delivery_sanitizer(canonical: list[str]):
    by_key = {re.sub(r"\s", "", c.lower()): c for c in canonical}

    def sanitize_delivery(v) -> str:
        if _is_blank(v):
            return INVALID
        key = re.sub(r"[^a-z]", "", str(v).lower())              # tira pontuação/hífen/espaço
        if key in by_key:
            return by_key[key]
        close = difflib.get_close_matches(key, by_key.keys(), n=1, cutoff=0.85)  # DELIVERD
        return by_key[close[0]] if close else INVALID

    return sanitize_delivery


# ----------------------------------------------------------------------------
# 4. Pipeline
# ----------------------------------------------------------------------------

def build_sanitizers(schema: dict) -> dict:
    return {
        "SaleDateSanitized": sanitize_date,
        "PorscheModelSanitized": make_model_sanitizer(schema["models"]),
        "ModelYearSanitized": sanitize_model_year,
        "SalesPriceSanitized": sanitize_price,
        "VehicleMileageSanitized": sanitize_mileage,
        "PayMethodSanitized": make_payment_sanitizer(schema["payments"]),
        "CitySanitized": sanitize_city,
        "StateSanitized": sanitize_state,
        "DeliveryStatusSanitized": make_delivery_sanitizer(schema["deliveries"]),
    }


def run_quality_checks(rows_out, out_cols, orig_cols, schema, src_headers) -> list[tuple[str, bool, str]]:
    """rows_out: lista de dicts {coluna: valor} já com as sanitizadas."""
    checks = []
    idx = {c: i for i, c in enumerate(out_cols)}

    checks.append(("Nenhuma coluna original removida",
                   all(c in out_cols for c in orig_cols), ""))

    adjacent = all(
        out_cols[idx[src] + 1] == out
        for src, out in schema["outputs"]
    )
    checks.append(("Colunas sanitizadas logo após a coluna de origem", adjacent, ""))

    def col(name):
        return [r[name] for r in rows_out]

    ok_dates = all(v == INVALID or re.fullmatch(r"\d{4}-\d{2}-\d{2}", v) for v in col("SaleDateSanitized"))
    checks.append(("Datas em ISO YYYY-MM-DD ou INVALID", ok_dates, ""))

    checks.append(("Estados com sigla USPS ou INVALID",
                   all(v == INVALID or v in _ABBR for v in col("StateSanitized")), ""))

    checks.append(("Milhagem em milhas inteiras",
                   all(v == INVALID or isinstance(v, int) for v in col("VehicleMileageSanitized")), ""))

    checks.append(("Preço numérico com duas casas decimais",
                   all(v == INVALID or (isinstance(v, float) and round(v, 2) == v)
                       for v in col("SalesPriceSanitized")), ""))

    checks.append(("Ano com 4 dígitos (1990–2035) ou INVALID",
                   all(v == INVALID or (isinstance(v, int) and 1990 <= v <= 2035)
                       for v in col("ModelYearSanitized")), ""))

    no_blank = all(not _is_blank(r[out]) for r in rows_out for _, out in schema["outputs"])
    checks.append(("Nenhum valor em branco nas colunas sanitizadas", no_blank, ""))

    deliv_ok = all(v in schema["deliveries"] or v == INVALID for v in col("DeliveryStatusSanitized"))
    checks.append(("Status de entrega dentro do vocabulário controlado", deliv_ok, ""))
    return checks


def main(src_path: str, schema_path: str, out_path: str):
    schema = load_schema(schema_path)
    sanitizers = build_sanitizers(schema)

    wb = openpyxl.load_workbook(src_path)
    ws = wb.active
    header = [c.value for c in ws[1]]

    missing = [c for c in schema["required"] if c not in header]
    if missing:
        raise SystemExit(f"Colunas obrigatórias ausentes na planilha: {missing}")

    out_for_src = dict(schema["outputs"])                         # source -> output
    out_cols: list[str] = []
    for h in header:
        out_cols.append(h)
        if h in out_for_src:
            out_cols.append(out_for_src[h])

    # --- constrói nova planilha preservando estilo das colunas originais -----
    new = openpyxl.Workbook()
    nws = new.active
    nws.title = ws.title
    san_fill = PatternFill("solid", fgColor="1F3864")
    new_vals_style = Font(name=ws["A2"].font.name, size=ws["A2"].font.sz)
    rows_out = []
    invalids = []   # (sale_id, coluna origem, valor bruto)

    col_pos = {c: i + 1 for i, c in enumerate(out_cols)}
    src_idx = {h: i + 1 for i, h in enumerate(header)}

    for r in range(1, ws.max_row + 1):
        record = {}
        for h in header:
            sc = ws.cell(r, src_idx[h])
            nc = nws.cell(r, col_pos[h], sc.value)
            if sc.has_style:
                nc.font, nc.fill, nc.border = copy(sc.font), copy(sc.fill), copy(sc.border)
                nc.alignment, nc.number_format = copy(sc.alignment), sc.number_format
            record[h] = sc.value
            if h in out_for_src:
                out_name = out_for_src[h]
                oc = nws.cell(r, col_pos[out_name])
                if r == 1:
                    oc.value = out_name
                    oc.font = copy(sc.font)
                    oc.fill = san_fill
                    oc.alignment = copy(sc.alignment)
                    oc.border = copy(sc.border)
                else:
                    val = sanitizers[out_name](sc.value)
                    oc.value = val
                    oc.font = new_vals_style
                    oc.alignment = Alignment(horizontal="left" if isinstance(val, str) else "right")
                    if out_name == "SalesPriceSanitized" and val != INVALID:
                        oc.number_format = "0.00"
                    record[out_name] = val
                    if val == INVALID:
                        invalids.append((record.get("sale_id"), h, sc.value, out_name))
        if r > 1:
            rows_out.append(record)

    # larguras de coluna
    for h in header:
        w = ws.column_dimensions[openpyxl.utils.get_column_letter(src_idx[h])].width or 15
        nws.column_dimensions[openpyxl.utils.get_column_letter(col_pos[h])].width = w
    for src, out in schema["outputs"]:
        nws.column_dimensions[openpyxl.utils.get_column_letter(col_pos[out])].width = max(len(out) + 4, 16)
    nws.freeze_panes = "A2"
    nws.auto_filter.ref = nws.dimensions

    # --- verificações de qualidade ------------------------------------------
    checks = run_quality_checks(rows_out, out_cols, header, schema, header)

    # --- aba de resumo -------------------------------------------------------
    sws = new.create_sheet("Resumo")
    bold = Font(name="Arial", bold=True)
    norm = Font(name="Arial")
    sws["A1"], sws["A1"].font = "Resumo da sanitização", Font(name="Arial", bold=True, size=13)
    sws.append([])
    sws.append(["Coluna sanitizada", "Origem", "Registros", "INVALID"])
    for c in sws[3]:
        c.font = bold
    for src, out in schema["outputs"]:
        n_inv = sum(1 for r in rows_out if r[out] == INVALID)
        sws.append([out, src, len(rows_out), n_inv])
    sws.append([])
    sws.append(["Verificações de qualidade", "Resultado"])
    for c in sws[sws.max_row]:
        c.font = bold
    for name, ok, _ in checks:
        sws.append([name, "OK" if ok else "FALHOU"])
    for row in sws.iter_rows(min_row=4):
        for c in row:
            if not c.font.b:
                c.font = norm
    sws.column_dimensions["A"].width = 52
    for L in "BCD":
        sws.column_dimensions[L].width = 14

    # --- aba com valores inválidos para revisão ------------------------------
    iws = new.create_sheet("Invalidos")
    iws.append(["sale_id", "coluna_origem", "valor_bruto", "coluna_sanitizada"])
    for c in iws[1]:
        c.font = bold
    for sid, h, raw, outn in invalids:
        iws.append([sid, h, str(raw), outn])
    for L, w in zip("ABCD", (10, 18, 30, 28)):
        iws.column_dimensions[L].width = w

    new.save(out_path)

    # --- relatório no console -----------------------------------------------
    print(f"Linhas processadas : {len(rows_out)}")
    print(f"Arquivo gerado     : {out_path}")
    print("INVALID por coluna :")
    for src, out in schema["outputs"]:
        print(f"  {out:<26} {sum(1 for r in rows_out if r[out] == INVALID)}")
    print("Verificações:")
    for name, ok, _ in checks:
        print(f"  [{'OK' if ok else 'FALHOU'}] {name}")
    if not all(ok for _, ok, _ in checks):
        raise SystemExit(1)


if __name__ == "__main__":
    a = sys.argv
    main(
        a[1] if len(a) > 1 else "/mnt/user-data/uploads/Tabela_Porche.xlsx",
        a[2] if len(a) > 2 else "/mnt/user-data/uploads/Schema_Porsche.md",
        a[3] if len(a) > 3 else "/mnt/user-data/outputs/Tabela_Porsche_Sanitizada.xlsx",
    )
